"""Secure file upload + text extraction for all supported file types."""
import importlib
import os
import io
import re
import hashlib
import zipfile
from typing import Optional
from config import UPLOAD_FOLDER, ALLOWED_EXTENSIONS

# Validation
def allowed_file(filename: str) -> bool:
    if "." not in filename:
        return False
    ext = filename.rsplit(".", 1)[1].lower()
    return ext in ALLOWED_EXTENSIONS
def safe_filename(filename: str) -> str:
    """Strip path components and dangerous characters."""
    base = os.path.basename(filename)
    base = re.sub(r"[^A-Za-z0-9._-]", "_", base)
    return base[:200]
def save_upload(file_storage, user_id: int):
    """Validate and store an uploaded file. Returns (stored_path, stored_name, ext, size)."""
    if not file_storage or not file_storage.filename:
        raise ValueError("No file provided.")
    if not allowed_file(file_storage.filename):
        raise ValueError("File type is not allowed.")
    filename = safe_filename(file_storage.filename)
    ext = filename.rsplit(".", 1)[1].lower()
    # Unique folder per user
    user_dir = os.path.join(UPLOAD_FOLDER, f"user_{user_id}")
    os.makedirs(user_dir, exist_ok=True)
    stored_name = f"{int(__import__('time').time())}_{filename}"
    stored_path = os.path.join(user_dir, stored_name)
    file_storage.save(stored_path)
    size = os.path.getsize(stored_path)
    return stored_path, stored_name, ext, size

# Text extraction
def extract_text(file_path: str, ext: str) -> str:
    ext = ext.lower()
    try:
        if ext == "txt":
            return _read_text(file_path)
        if ext == "csv":
            return _read_text(file_path)
        if ext in ("png", "jpg", "jpeg"):
            return _extract_image_ocr(file_path, ext)
        if ext == "pdf":
            return _extract_pdf(file_path)
        if ext in ("doc", "rtf", "odt"):
            return _extract_doc(file_path, ext)
        if ext == "docx":
            return _extract_docx(file_path)
        if ext == "pptx":
            return _extract_pptx(file_path)
        if ext == "ppt":
            return _extract_ppt_legacy(file_path)
        if ext in ("xls", "xlsx"):
            return _extract_spreadsheet(file_path, ext)
    except Exception as exc:
        return f"[Extraction error: {exc}]"
    return ""


def extraction_failed(text: str) -> bool:
    """Return true for extraction markers that must never be scanned."""
    value = (text or "").strip().lower()
    return value.startswith("[") and (
        "extraction error" in value
        or "ocr not available" in value
        or "text extraction skipped" in value
    )
def _read_text(path: str) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as fh:
                return fh.read()
        except UnicodeDecodeError:
            continue
    return ""
def _extract_pdf(path: str) -> str:
    text_parts = []
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                t = page.extract_text() or ""
                text_parts.append(t)
    except Exception:
        pass
    if not "".join(text_parts).strip():
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(path)
            for page in reader.pages:
                text_parts.append(page.extract_text() or "")
        except Exception:
            pass
    return "\n".join(text_parts).strip()

def _extract_docx(path: str) -> str:
    try:
        from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception as exc:
        return f"[DOCX extraction error: {exc}]"
def _extract_doc(path: str, ext: str) -> str:
    if ext == "rtf":
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as fh:
                raw = fh.read()
            # Strip RTF control words
            plain = re.sub(r"\\[a-zA-Z]+\d* ?|[\{\}]", "", raw)
            plain = re.sub(r"\\'([0-9a-fA-F]{2})", ".", plain)
            return plain
        except Exception as exc:
            return f"[RTF extraction error: {exc}]"
    # Legacy .doc / .odt - try as zip (odt) or fall back to raw text scan
    if ext == "odt":
        try:
            with zipfile.ZipFile(path) as zf:
                with zf.open("content.xml") as f:
                    xml = f.read().decode("utf-8", errors="ignore")
                    return re.sub(r"<[^>]+>", " ", xml)
        except Exception as exc:
            return f"[ODT extraction error: {exc}]"
    # .doc binary - best effort raw text scan
    try:
        with open(path, "rb") as fh:
            raw = fh.read()
        text = raw.decode("latin-1", errors="ignore")
        return re.sub(r"[^\x20-\x7E\n]", " ", text)
    except Exception as exc:
        return f"[DOC extraction error: {exc}]"

def _extract_pptx(path: str) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(path)
        out = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    out.append(shape.text_frame.text)
        return "\n".join(out)
    except Exception as exc:
        return f"[PPTX extraction error: {exc}]"

def _extract_ppt_legacy(path: str) -> str:
    try:
        with open(path, "rb") as fh:
            raw = fh.read()
        text = raw.decode("latin-1", errors="ignore")
        return re.sub(r"[^\x20-\x7E\n]", " ", text)
    except Exception as exc:
        return f"[PPT extraction error: {exc}]"

def _extract_spreadsheet(path: str, ext: str) -> str:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                out.append(" ".join(str(c) for c in row if c is not None))
        return "\n".join(out)
    except Exception as exc:
        return f"[Spreadsheet extraction error: {exc}]"

def _extract_image_ocr(path: str, ext: str) -> str:
    """Images: attempt OCR with Tesseract if available, else return a marker."""
    try:
        pytesseract = importlib.import_module("pytesseract")
        Image = importlib.import_module("PIL.Image")
        return pytesseract.image_to_string(Image.open(path))
    except Exception:
        return "[Image uploaded - OCR not available in this environment. Text extraction skipped.]"
def text_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()